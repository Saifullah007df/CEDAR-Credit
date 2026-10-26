"""Self-hosted v6 additions: signed-off journal imports, offline CSV/XLSX, policies and evidence review.
Does not connect to any third-party API or modify an external lender ledger.
"""
import csv, io, json, hashlib
from datetime import date
from decimal import Decimal
from typing import List
from fastapi import APIRouter, Depends, File, UploadFile, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from backend.app import (
    db, require, user_from_header, money, ds, uid, D, Loan, Payment, Invoice, Doc, Journal,
    JournalLine, PolicyVersion, DocumentCorrection, Audit, Batch, SessionLocal,
    BatchIn, PaymentIn, ingest, need_loan, ensure_loan_scope, LoanAccess, touch, audit, outbox, active_categories
)
router=APIRouter(prefix='/v1')

class PostingLine(BaseModel):
    account: str=Field(min_length=2,max_length=80)
    debit:Decimal=D('0')
    credit:Decimal=D('0')
class JournalIn(BaseModel):
    source:str=Field(min_length=2,max_length=80)
    external_id:str=Field(min_length=1,max_length=128)
    loan_id:str
    posting_date:str
    journal_type:str='OTHER'
    currency:str='INR'
    lines:List[PostingLine]=Field(min_length=2,max_length=200)

@router.post('/journals',status_code=201)
def create_journal(body:JournalIn,s:Session=Depends(db),u=Depends(require('admin','operator'))):
    loan=need_loan(s,body.loan_id);ensure_loan_scope(s,u,body.loan_id)
    try: date.fromisoformat(body.posting_date)
    except ValueError: raise HTTPException(422,'posting_date must use YYYY-MM-DD')
    if body.currency!='INR':raise HTTPException(422,'Only INR supported in this release')
    if s.scalar(select(Journal).where(Journal.source==body.source, Journal.external_id==body.external_id)):
        raise HTTPException(409,'Duplicate journal source identity; changes require controlled revision')
    debit=D('0');credit=D('0')
    for line in body.lines:
        if line.debit < 0 or line.credit < 0 or (line.debit > 0 and line.credit > 0) or (line.debit==0 and line.credit==0):
            raise HTTPException(422,'Each line must contain one strictly positive debit or credit')
        value=line.debit if line.debit>0 else line.credit
        money(value)
        debit+=line.debit;credit+=line.credit
    if debit!=credit or debit<=0:raise HTTPException(422,'Journal debits must equal credits and be positive')
    j=Journal(source=body.source,external_id=body.external_id,loan_id=body.loan_id,
        posting_date=body.posting_date,journal_type=body.journal_type.upper(),currency=body.currency)
    s.add(j);s.flush()
    for line in body.lines:
        s.add(JournalLine(journal_id=j.id,account=line.account.upper(),debit=line.debit,credit=line.credit))
    touch(s,loan)
    audit(s,u.username,'JOURNAL_IMPORTED',j.id,{'debits':ds(debit),'credits':ds(credit),'source':body.source})
    outbox(s,'JournalImported',j.id,{'loan_id':loan.id})
    s.commit()
    return {'id':j.id,'status':'BALANCED','debit':ds(debit),'credit':ds(credit)}

@router.get('/journals')
def journals(loan_id:str|None=None,s:Session=Depends(db),u=Depends(require('admin','operator','reviewer','approver','auditor','lender'))):
    stmt=select(Journal).order_by(Journal.posting_date)
    if loan_id:stmt=stmt.where(Journal.loan_id==loan_id)
    from backend.app import DEMO_MODE
    if not DEMO_MODE and u.role!='admin':stmt=stmt.where(Journal.loan_id.in_(select(LoanAccess.loan_id).where(LoanAccess.user_id==u.id)))
    arr=[]
    for j in s.scalars(stmt).all():
        lines=s.scalars(select(JournalLine).where(JournalLine.journal_id==j.id)).all()
        arr.append({'id':j.id,'source':j.source,'external_id':j.external_id,'loan_id':j.loan_id,
            'posting_date':j.posting_date,'journal_type':j.journal_type,
            'debit':ds(sum((x.debit for x in lines),D('0'))),
            'credit':ds(sum((x.credit for x in lines),D('0'))),
            'lines':[{'account':x.account,'debit':ds(x.debit),'credit':ds(x.credit)} for x in lines]})
    return arr

@router.get('/loans/{loan_id}/ledger-reconciliation')
def reconcile_ledger(loan_id:str,s:Session=Depends(db),u=Depends(require('admin','operator','reviewer','approver','auditor','lender'))):
    loan=need_loan(s,loan_id);ensure_loan_scope(s,u,loan_id)
    js=s.scalars(select(Journal).where(Journal.loan_id==loan_id)).all()
    journal_ids=[j.id for j in js]
    lines=s.scalars(select(JournalLine).where(JournalLine.journal_id.in_(journal_ids))).all() if journal_ids else []
    debit=sum((x.debit for x in lines),D('0'));credit=sum((x.credit for x in lines),D('0'))
    disbursement_ids=[j.id for j in js if j.journal_type=='DISBURSEMENT']
    receivable=s.scalars(select(JournalLine).where(JournalLine.journal_id.in_(disbursement_ids),JournalLine.account=='LOAN_RECEIVABLE')).all() if disbursement_ids else []
    disbursed=sum((x.debit-x.credit for x in receivable),D('0'))
    findings=[]
    if not js:findings.append('No authoritative journal records imported; financial reconciliation unavailable')
    if debit!=credit:findings.append('Unbalanced ledger postings detected')
    if not disbursement_ids:findings.append('No disbursement journal observed; principal cannot be reconciled')
    if disbursement_ids and disbursed!=loan.principal:findings.append('Observed disbursement differs from contractual principal; check tranches and source coverage')
    incomplete_batches=s.scalar(select(func.count()).select_from(Batch).join(Payment,Payment.batch_id==Batch.id).where(Payment.loan_id==loan_id,Batch.status!='COMPLETE')) or 0
    if incomplete_batches:findings.append('At least one financial import batch is incomplete')
    return {'loan_id':loan_id,'status':'PROVISIONAL' if findings else 'RECONCILED_WITH_IMPORTED_JOURNALS',
        'journal_count':len(js),'debit_total':ds(debit),'credit_total':ds(credit),'observed_disbursement':ds(disbursed),
        'contractual_principal':ds(loan.principal),'source_batches_incomplete':incomplete_batches>0,
        'findings':findings,'caveat':'Only imported source records are checked; completeness against the external LMS is not independently established.'}

@router.post('/imports/financial-file',status_code=201)
async def import_financial_file(source:str,batch_ref:str,expected_count:int,
    file:UploadFile=File(...),s:Session=Depends(db),u=Depends(require('admin','operator'))):
    data=await file.read(5_000_001)
    if len(data)>5_000_000:raise HTTPException(413,'File exceeds 5MB cap')
    name=(file.filename or '').lower()
    try:
        if name.endswith('.csv'):
            content=data.decode('utf-8-sig')
            rows=list(csv.DictReader(io.StringIO(content)))
        elif name.endswith('.xlsx'):
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            ws=wb.active
            iterator=ws.iter_rows(values_only=True)
            header=[str(v or '').strip() for v in next(iterator)]
            rows=[dict(zip(header, values)) for values in iterator]
            wb.close()
        else:raise HTTPException(415,'Use UTF-8 CSV or XLSX (.xls legacy not supported)')
        if len(rows)>3000:raise HTTPException(413,'Max 3000 transactions per file')
        required={'loan_id','external_id','payee','amount'}
        if not rows or not required.issubset(rows[0]):raise HTTPException(422,'Required columns: loan_id,external_id,payee,amount')
        parsed=[PaymentIn(loan_id=str(r['loan_id']).strip(),external_id=str(r['external_id']).strip(),
            payee=str(r['payee']).strip(),amount=D(str(r['amount']))) for r in rows]
        if any(not v.loan_id or not v.external_id or not v.payee for v in parsed):
            raise HTTPException(422,'Blank required value')
    except HTTPException:raise
    except Exception as ex:raise HTTPException(422,f'Invalid financial file format: {type(ex).__name__}')
    return ingest(BatchIn(source=source,batch_ref=batch_ref,expected_count=expected_count,payments=parsed),s,u)

class PolicyIn(BaseModel):
    categories:list[str]=Field(min_length=1,max_length=50)
    effective_date:str
    reason:str=Field(min_length=5)

@router.post('/loans/{loan_id}/policies',status_code=201)
def publish_policy(loan_id:str,body:PolicyIn,s:Session=Depends(db),u=Depends(require('admin','approver'))):
    loan=need_loan(s,loan_id);ensure_loan_scope(s,u,loan_id)
    try:date.fromisoformat(body.effective_date)
    except ValueError:raise HTTPException(422,'Invalid effective date')
    current=s.scalar(select(PolicyVersion).where(PolicyVersion.loan_id==loan_id).order_by(PolicyVersion.version.desc()))
    if current and body.effective_date<current.effective_date:
        raise HTTPException(409,'Backdated policy amendments require a separate authorized workflow')
    version=(current.version+1) if current else 1
    categories=sorted(set(x.strip().upper() for x in body.categories if x.strip()))
    if not categories:raise HTTPException(422,'At least one category required')
    pv=PolicyVersion(loan_id=loan_id,version=version,categories_json=json.dumps(categories),
        effective_date=body.effective_date,approved_by=u.username)
    s.add(pv);touch(s,loan)
    audit(s,u.username,'PURPOSE_POLICY_PUBLISHED',loan_id,{'version':version,'categories':categories,'reason':body.reason})
    outbox(s,'PolicyPublished',loan_id,{'loan_id':loan_id,'version':version})
    s.commit()
    return {'loan_id':loan_id,'version':version,'categories':categories,'approved_by':u.username}

@router.get('/loans/{loan_id}/policies')
def policies(loan_id:str,s:Session=Depends(db),u=Depends(require('admin','operator','reviewer','approver','auditor','lender'))):
    need_loan(s,loan_id);ensure_loan_scope(s,u,loan_id)
    return [{'version':p.version,'effective_date':p.effective_date,
        'categories':json.loads(p.categories_json),'approved_by':p.approved_by}
        for p in s.scalars(select(PolicyVersion).where(PolicyVersion.loan_id==loan_id).order_by(PolicyVersion.version)).all()]

class DocCorrectionIn(BaseModel):
    expected_version:int
    corrected_text:str=Field(min_length=1,max_length=10000)
    reason:str=Field(min_length=5,max_length=500)

@router.post('/documents/{doc_id}/corrections')
def correct_document(doc_id:str,body:DocCorrectionIn,s:Session=Depends(db),u=Depends(require('reviewer','admin'))):
    d=s.get(Doc,doc_id)
    if not d:raise HTTPException(404,'Document not found')
    ensure_loan_scope(s,u,d.loan_id)
    if d.version!=body.expected_version:raise HTTPException(409,'Stale document version')
    d.version+=1
    s.add(DocumentCorrection(doc_id=d.id,version=d.version,corrected_text=body.corrected_text,
        reason=body.reason,actor=u.username))
    touch(s,need_loan(s,d.loan_id))
    audit(s,u.username,'DOCUMENT_CORRECTED',d.id,{'version':d.version,'reason':body.reason})
    outbox(s,'EvidenceCorrected',d.id,{'loan_id':d.loan_id,'version':d.version})
    s.commit()
    return {'id':d.id,'version':d.version,'original_ocr_unchanged':True}

@router.get('/documents/{doc_id}/history')
def document_history(doc_id:str,s:Session=Depends(db),u=Depends(require('admin','operator','reviewer','approver','auditor'))):
    d=s.get(Doc,doc_id)
    if not d:raise HTTPException(404,'Document not found')
    ensure_loan_scope(s,u,d.loan_id)
    corrections=s.scalars(select(DocumentCorrection).where(DocumentCorrection.doc_id==doc_id).order_by(DocumentCorrection.version)).all()
    return {'id':d.id,'source_sha256':d.sha256,'original_ocr':d.extracted_text,
        'current_version':d.version,'corrections':[{'version':c.version,'corrected_text':c.corrected_text,
        'reason':c.reason,'actor':c.actor} for c in corrections]}

class CreateStaffIn(BaseModel):
    username:str=Field(min_length=3,max_length=80,pattern=r'^[a-zA-Z0-9._-]+$')
    password:str=Field(min_length=16,max_length=256)
    role:str

@router.post('/staff',status_code=201)
def create_staff(body:CreateStaffIn,s:Session=Depends(db),u=Depends(require('admin'))):
    from backend.app import User,ph
    if body.role not in {'admin','operator','reviewer','approver','auditor','lender'}:
        raise HTTPException(422,'Invalid staff role')
    if len(set(body.password))<8:raise HTTPException(422,'Password lacks sufficient character variety')
    if s.scalar(select(User).where(User.username==body.username)):
        raise HTTPException(409,'Username already in use')
    import secrets,base64
    from backend.app import DEMO_MODE
    secret=base64.b32encode(secrets.token_bytes(20)).decode('ascii') if not DEMO_MODE else None
    user=User(username=body.username,role=body.role,password_hash=ph.hash(body.password),totp_secret=secret)
    s.add(user);s.flush()
    audit(s,u.username,'STAFF_CREATED',user.id,{'username':body.username,'role':body.role})
    s.commit()
    response={'id':user.id,'username':body.username,'role':body.role}
    if secret:response['one_time_mfa_enrollment_secret']=secret
    return response

@router.get('/operations/status')
def operations_status(s:Session=Depends(db),u=Depends(require('admin','operator','auditor'))):
    from backend.app import Outbox
    pending=s.scalar(select(func.count()).select_from(Outbox).where(Outbox.published==False)) or 0
    failed=s.scalar(select(func.count()).select_from(Outbox).where(Outbox.published==False,Outbox.attempts>=5)) or 0
    return {'pending_outbox_events':pending,'events_with_five_or_more_attempts':failed,
        'offline_runtime':True,'outbound_financial_apis':False,
        'warning':'Source-record completeness against external institutions is not established without authorized source controls.'}


class GrantLoanIn(BaseModel):
    username:str
    loan_id:str

@router.post('/staff/loan-access',status_code=201)
def grant_loan_access(body:GrantLoanIn,s:Session=Depends(db),u=Depends(require('admin'))):
    from backend.app import User
    target=s.scalar(select(User).where(User.username==body.username))
    if not target:raise HTTPException(404,'Staff user not found')
    need_loan(s,body.loan_id)
    if s.get(LoanAccess,{'user_id':target.id,'loan_id':body.loan_id}):
        raise HTTPException(409,'Access already assigned')
    s.add(LoanAccess(user_id=target.id,loan_id=body.loan_id))
    audit(s,u.username,'LOAN_ACCESS_GRANTED',body.loan_id,{'staff_user':body.username})
    s.commit()
    return {'username':body.username,'loan_id':body.loan_id,'assigned':True}


@router.post('/staff/{user_id}/revoke-sessions')
def revoke_staff_sessions(user_id:str,s:Session=Depends(db),u=Depends(require('admin'))):
    from backend.app import User
    target=s.get(User,user_id)
    if not target:raise HTTPException(404,'Staff not found')
    target.token_version=(target.token_version or 0)+1
    audit(s,u.username,'STAFF_SESSIONS_REVOKED',user_id,{})
    s.commit()
    return {'id':user_id,'sessions_revoked':True}
